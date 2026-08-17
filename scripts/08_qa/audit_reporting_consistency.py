from __future__ import annotations

from uscanopy.paths import (analysis_store_root, external_input_root, manuscript_root, model_store_root, repo_root)

import csv
import json
import re
from pathlib import Path

import pandas as pd
from docx import Document


OUT = (
    analysis_store_root()
    / "Result"
    / "reproducibility_qa"
    / "reporting_audit"
)
MANUSCRIPT_DIR = Path(str(manuscript_root()))
RESULT5 = Path(str(analysis_store_root().joinpath('Result', 'Result5', 'versions', 'step2_multicriteria_pareto')))
TRADEOFF = Path(str(analysis_store_root().joinpath('Result', 'Result5', 'versions', 'step01_city_tradeoff_outcomes')))
STRUCTURAL = Path(str(analysis_store_root().joinpath('Result', 'Result5', 'versions', 'step02_divergence_decomposition')))


SEARCH_PATTERNS = {
    "version_label": re.compile(r"\bV(?:ersion\s*)?(?:3[1-9]|4[0-3])\b|review package", re.I),
    "raw_missing": re.compile(r"\bnan\b|not available", re.I),
    "independence_wording": re.compile(r"independent (?:axes|objective|criteria)|independently defined", re.I),
    "allocation_claim": re.compile(r"plant(?:ing|able)|benefit|people reached|resident reach|implementation", re.I),
    "budget_wording": re.compile(r"equal 10% budget|same 10% budget|10% budget", re.I),
    "result5_values": re.compile(r"0\.9(?:1|2|3|4|38|40)|0\.7(?:09|12)|1\.(?:09|11|13|14|18|21|27)|0\.(?:78|87)"),
}


def all_doc_text(doc: Document):
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            yield "paragraph", str(index), text
    for table_index, table in enumerate(doc.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            for col_index, cell in enumerate(row.cells, start=1):
                text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                if text:
                    yield "table_cell", f"{table_index}:{row_index}:{col_index}", text


def find_documents() -> list[Path]:
    expected = [
        MANUSCRIPT_DIR / "manuscript.docx",
        MANUSCRIPT_DIR / "supplementary_information.docx",
    ]
    missing = [str(path) for path in expected if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing reporting-audit documents: " + "; ".join(missing))
    return expected


def result5_source_of_truth() -> dict:
    intervals = pd.read_csv(RESULT5 / "tables" / "04_two_level_rule_intervals.csv")
    groups = pd.read_csv(RESULT5 / "tables" / "05_two_level_group_intervals.csv")

    selected_rules = [
        "population_oracle",
        "heat_context_oracle",
        "social_condition_oracle",
        "pareto_0.50",
    ]
    metric_records = []
    for _, row in intervals.loc[intervals["rule"].isin(selected_rules)].iterrows():
        metric_records.append(
            {
                "rule": row["rule"],
                "metric": row["metric"],
                "point_estimate": float(row["equal_city_central"]),
                "bootstrap_median_diagnostic_only": float(row["two_level_median"]),
                "bootstrap_ci_low": float(row["two_level_ci_low"]),
                "bootstrap_ci_high": float(row["two_level_ci_high"]),
            }
        )

    group_records = []
    for _, row in groups.loc[groups["rule"].eq("pareto_0.50")].iterrows():
        group_records.append(
            {
                "rule": row["rule"],
                "group": row["group"],
                "point_estimate": float(row["equal_city_central"]),
                "bootstrap_median_diagnostic_only": float(row["two_level_median"]),
                "bootstrap_ci_low": float(row["two_level_ci_low"]),
                "bootstrap_ci_high": float(row["two_level_ci_high"]),
            }
        )

    headline = pd.read_csv(
        TRADEOFF / "tables" / "national_primary_tradeoff_outcomes.csv"
    )
    structural = pd.read_csv(
        STRUCTURAL / "tables" / "structural_finding_evidence_ledger.csv"
    )

    return {
        "reporting_rule": {
            "point_estimate": "Mean of the 50 city-specific deterministic full-frame estimates (equal city weight).",
            "uncertainty": "2.5th and 97.5th percentiles from the two-level city and within-city spatial-block bootstrap.",
            "bootstrap_median": "Diagnostic only; never substitute for the deterministic point estimate.",
            "missing_values": "Render as an em dash and define as not estimated or not applicable; never print nan.",
        },
        "metric_records": metric_records,
        "group_records": group_records,
        "headline_records": headline.to_dict(orient="records"),
        "structural_records": structural.to_dict(orient="records"),
    }


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    rows = []
    for path in find_documents():
        document = Document(path)
        for location_type, location, text in all_doc_text(document):
            matched = [name for name, pattern in SEARCH_PATTERNS.items() if pattern.search(text)]
            if matched:
                rows.append(
                    {
                        "document": path.name,
                        "location_type": location_type,
                        "location": location,
                        "issue_classes": ";".join(matched),
                        "text": text,
                    }
                )

    with (OUT / "reporting_occurrences.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=["document", "location_type", "location", "issue_classes", "text"])
        writer.writeheader()
        writer.writerows(rows)

    source = {
        "manuscript_target": "submitted analysis release",
        "analysis_archive_note": "Role-based analytical folders preserve the frozen dependency graph for the submitted release.",
        "result1": {
            "analysis_units": 38618,
            "positive_population_units": 37721,
            "cities": 50,
            "represented_2024_acs_population": 50477302,
            "resident_context_min_pct": 3.3697,
            "resident_context_min_city": "El Paso",
            "resident_context_max_pct": 41.6804,
            "resident_context_max_city": "Raleigh",
            "primary_median_resident_minus_area_pp": 0.68076,
            "primary_cities_positive": 29,
            "fully_contained_median_resident_minus_area_pp": -0.15518,
            "fully_contained_cities_positive": 23,
        },
        "result5": result5_source_of_truth(),
        "terminology": {
            "objectives": "distinct, separately defined objectives",
            "allocation": "diagnostic equal-budget allocation experiment",
            "population_metric": "resident-alignment efficiency",
            "budget": "10% of each city's analytical canopy-deficit capacity",
            "pareto_half": "illustrative equal-weight compromise (lambda = 0.5)",
            "shade_boundary": "visible woody canopy context, not cast shade",
        },
    }
    (OUT / "single_source_of_truth.json").write_text(
        json.dumps(source, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    documents = find_documents()
    main_text = "\n".join(text for _, _, text in all_doc_text(Document(documents[0])))
    supplement_text = "\n".join(text for _, _, text in all_doc_text(Document(documents[1])))
    required_main = ("0.355", "47.7%", "28.8%", "22.9%")
    required_supplement = (
        "Supplementary Table 7",
        "Controlled structural decomposition",
    )
    old_version_hits = [
        row for row in rows if "version_label" in row["issue_classes"].split(";")
    ]
    raw_missing_hits = [
        row for row in rows if "raw_missing" in row["issue_classes"].split(";")
    ]
    checks = {
        "main_tradeoff_values_present": all(token in main_text for token in required_main),
        "supplement_structural_evidence_present": all(
            token in supplement_text for token in required_supplement
        ),
        "no_stale_version_labels": not old_version_hits,
        "no_raw_missing_tokens": not raw_missing_hits,
        "structural_ledger_has_nine_rows": len(
            source["result5"]["structural_records"]
        ) == 9,
    }
    failures = [name for name, passed in checks.items() if not passed]
    audit = {
        "status": "PASS" if not failures else "FAIL",
        "documents_scanned": 2,
        "flagged_occurrences": len(rows),
        "checks": checks,
        "failures": failures,
        "source_tables": [
            str(RESULT5 / "tables" / "04_two_level_rule_intervals.csv"),
            str(RESULT5 / "tables" / "05_two_level_group_intervals.csv"),
            str(TRADEOFF / "tables" / "national_primary_tradeoff_outcomes.csv"),
            str(STRUCTURAL / "tables" / "structural_finding_evidence_ledger.csv"),
        ],
        "frozen_rule": "Deterministic equal-city mean for point estimates; two-level bootstrap percentiles for intervals.",
    }
    (OUT / "reporting_audit.json").write_text(
        json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps(audit, indent=2, ensure_ascii=False))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
